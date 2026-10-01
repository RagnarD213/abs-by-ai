#!/usr/bin/env python3
"""Abs By AI clip library: one catalog of every reusable AI clip and B-roll clip.

Search it BEFORE generating a new AI clip, buying/searching stock, or hunting raw rolls.
Register every newly approved clip with `add` so the library keeps growing.

  find "man eating salad" [--aspect 9x16] [--kind ai] [--people none] [--category food-nutrition] [--rolls]
  show A0042
  add FILE --kind ai --slug man-eating-salad-kitchen --description "..." --category food-nutrition
      [--people other-man] [--model kling-v3] [--prompt "..."] [--used-in "RO-05"] [--status used-final]
  ingest CANDIDATES.json [...]        (bulk build from sweep files; --dry-run first)
  thumbs | describe | sheet | drive-sync | verify

Catalog (source of truth): Media/clip-library/catalog.json  (gitignored; repo is public)
Files: /Volumes/Extreme/_asset_library_stage/Abs By AI - Video Asset Library/{03,04}/<category>/
Drive mirror: folder 1Hby8O4mB4HZS341qvrVKHSHyCGBgP8mi (anyone-with-link, children inherit)
"""

import argparse
import base64
import datetime as dt
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "rolls"))
import roll_sidecar as rs  # noqa: E402  (probe, content_identity, binary, secret_value)

REPO = HERE.parents[3]
CAT_DIR = Path(os.environ.get("CLIPLIB_CATALOG_DIR", REPO / "Media" / "clip-library"))
CATALOG = CAT_DIR / "catalog.json"
THUMBS = CAT_DIR / "thumbs"
CONTACT = CAT_DIR / "contact"
LIB = Path(os.environ.get("CLIPLIB_ROOT", "/Volumes/Extreme/_asset_library_stage/Abs By AI - Video Asset Library"))
DRIVE_ROOT = "1Hby8O4mB4HZS341qvrVKHSHyCGBgP8mi"
RCLONE = str(Path.home() / "bin" / "rclone")
REJECTS = Path("/Volumes/Extreme/_edit_work/_clip_library_set_aside")

AI_DIR = "04 AI-Generated Clips"
BROLL_DIR = "03 B-Roll - Real Footage"
AI_CATEGORIES = ["ai-dan", "people", "food-nutrition", "gym-and-exercise", "lifestyle",
                 "concepts-and-gags", "exercise-demos"]
BROLL_CATEGORIES = ["dan-filmed", "screen-recordings", "stock"]
STOCK_SUBS = ["food-nutrition", "gym-and-exercise", "lifestyle", "people"]
KINDS = ["ai", "real", "screen", "stock"]
STATUSES = ["used-final", "usable-unused"]
PEOPLE = ["dan", "ai-dan", "other-man", "woman", "mixed", "none"]


def now():
    return dt.datetime.now().astimezone().isoformat(timespec="seconds")


def load():
    if CATALOG.exists():
        return json.loads(CATALOG.read_text())
    return {"schema_version": 1, "library_root": str(LIB), "drive_root": DRIVE_ROOT, "clips": []}


def save(cat):
    CAT_DIR.mkdir(parents=True, exist_ok=True)
    cat["updated"] = now()
    cat["clips"].sort(key=lambda c: c["id"])
    tmp = CATALOG.with_suffix(".tmp")
    tmp.write_text(json.dumps(cat, indent=1, ensure_ascii=False) + "\n")
    tmp.replace(CATALOG)


def slugify(text, limit=60):
    s = re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-")
    return s[:limit].rstrip("-") or "clip"


def aspect_label(w, h):
    if not w or not h:
        return "unknown"
    r = w / h
    for label, val in (("9x16", 9 / 16), ("16x9", 16 / 9), ("1x1", 1.0), ("4x5", 0.8), ("2x3", 2 / 3), ("3x4", 0.75), ("4x3", 4 / 3)):
        if abs(r - val) < 0.03:
            return label
    return "{}x{}".format(w, h)


def folder_for(kind, category):
    if kind == "ai":
        return Path(AI_DIR) / (category if category in AI_CATEGORIES else "concepts-and-gags")
    if kind == "screen":
        return Path(BROLL_DIR) / "screen-recordings"
    if kind == "stock":
        return Path(BROLL_DIR) / "stock" / (category if category in STOCK_SUBS else "other")
    return Path(BROLL_DIR) / "dan-filmed"


def next_id(cat, kind):
    prefix = "A" if kind == "ai" else "B"
    nums = [int(c["id"][1:]) for c in cat["clips"] if c["id"].startswith(prefix)]
    return "{}{:04d}".format(prefix, (max(nums) + 1) if nums else 1)


def by_key(cat):
    return {c["content_key"]: c for c in cat["clips"]}


def probe_file(path):
    p = rs.probe(path)
    w, h = [int(x) for x in p["picture"]["decoded_resolution"].split("x")]
    return {"duration": p["duration"], "width": w, "height": h, "fps": p["picture"]["fps"],
            "codec": p["picture"]["codec"], "has_audio": bool(p["audio_streams"]),
            "aspect": aspect_label(w, h)}


def make_thumbs(clip):
    """Poster frame (480 px) + 6-frame contact strip, kept on the Mac so search works offline."""
    src = LIB / clip["library_path"]
    THUMBS.mkdir(parents=True, exist_ok=True)
    CONTACT.mkdir(parents=True, exist_ok=True)
    ff = rs.binary("ffmpeg")
    d = max(clip["duration"], 0.2)
    poster = THUMBS / (clip["id"] + ".jpg")
    subprocess.run([ff, "-v", "error", "-y", "-ss", "{:.2f}".format(d * 0.4), "-i", str(src), "-frames:v", "1",
                    "-vf", "scale=480:-2", "-q:v", "4", str(poster)], check=True)
    sheet = CONTACT / (clip["id"] + ".jpg")
    fps = 6.0 / d
    cols, tile_w = (6, 200) if clip["height"] > clip["width"] else (3, 320)
    rows = 6 // cols
    subprocess.run([ff, "-v", "error", "-y", "-i", str(src), "-vf",
                    "fps={:.5f},scale={}:-2,tile={}x{}:padding=4:color=white".format(fps, tile_w, cols, rows),
                    "-frames:v", "1", "-q:v", "4", str(sheet)], check=True)


def place(src, clip, move):
    dest = LIB / clip["library_path"]
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        return
    if move:
        shutil.move(str(src), str(dest))
    else:
        shutil.copyfile(str(src), str(dest))  # copy2/copystat fails on the exFAT Extreme drive (chflags)
        try:
            os.utime(str(dest), (src.stat().st_atime, src.stat().st_mtime))
        except OSError:
            pass


def new_record(cat, src, kind, slug, fields):
    info = probe_file(src)
    ident = rs.content_identity(src)
    cid = next_id(cat, kind)
    name = "{}_{}_{}_{}s{}".format(cid, slugify(slug), info["aspect"], max(1, round(info["duration"])), src.suffix.lower())
    category = fields.get("category") or ""
    folder = folder_for(kind, category)
    if kind == "ai" and category not in AI_CATEGORIES:
        category = "concepts-and-gags"
    if kind in ("real", "screen"):
        category = folder.name
    if kind == "stock":
        category = "stock-" + folder.name
    rec = {
        "id": cid, "kind": kind, "category": category, "slug": slugify(slug),
        "file_name": name, "library_path": str(folder / name),
        "description": fields.get("description", ""), "tags": fields.get("tags", []),
        "people": fields.get("people") or "none", "status": fields.get("status") or "usable-unused",
        "used_in": fields.get("used_in", []), "needs_ai_label": kind == "ai" and (fields.get("people") or "none") != "none",
        "model": fields.get("model"), "prompt": fields.get("prompt"),
        "source_path": str(src), "alt_paths": fields.get("alt_paths", []), "legacy_name": src.name,
        "evidence": fields.get("evidence"), "notes": fields.get("notes", ""),
        "content_key": ident["content_key"], "size_bytes": ident["size_bytes"],
        "drive_id": None, "added": now(),
    }
    rec.update(info)
    return rec


# ---------------------------------------------------------------- commands

def cmd_ingest(args):
    cat = load()
    keys = by_key(cat)
    plan, skipped = [], []
    for f in args.files:
        area = json.loads(Path(f).read_text())
        for e in area.get("entries", []):
            p = Path(e.get("path", ""))
            if e.get("status") == "rejected":
                skipped.append(("rejected", str(p)))
                continue
            if not p.exists():
                skipped.append(("missing", str(p)))
                continue
            plan.append((p, e, area.get("area")))
    added = dup = 0
    for p, e, area in plan:
        ident = rs.content_identity(p)
        if ident["content_key"] in keys:
            rec = keys[ident["content_key"]]
            for u in e.get("used_in", []):
                if u not in rec["used_in"]:
                    rec["used_in"].append(u)
            if e.get("status") == "used-final":
                rec["status"] = "used-final"
            dup += 1
            continue
        kind = e.get("kind") if e.get("kind") in KINDS else "ai"
        if args.dry_run:
            print("ADD {:6} {:18} {}".format(kind, e.get("category", ""), e.get("slug")))
            keys[ident["content_key"]] = {"used_in": [], "status": ""}
            added += 1
            continue
        rec = new_record(cat, p, kind, e.get("slug") or p.stem, e)
        rec["sweep_area"] = area
        in_library = str(p).startswith(str(LIB / AI_DIR)) or str(p).startswith(str(LIB / BROLL_DIR))
        place(p, rec, move=in_library)
        try:
            make_thumbs(rec)
        except subprocess.CalledProcessError as exc:
            rec["notes"] = (rec["notes"] + " | thumbnail failed: {}".format(exc)).strip(" |")
        cat["clips"].append(rec)
        keys[rec["content_key"]] = rec
        added += 1
        save(cat)  # every record: a crash must never leave a moved file without its catalog row
        if added % 25 == 0:
            print("... {} added".format(added), flush=True)
    if not args.dry_run:
        save(cat)
    print("added {}, merged duplicates {}, skipped {}".format(added, dup, len(skipped)))
    for why, p in skipped:
        if why == "missing":
            print("  MISSING", p)


def cmd_add(args):
    cat = load()
    src = Path(args.file).expanduser().resolve()
    ident = rs.content_identity(src)
    for c in cat["clips"]:
        if c["content_key"] == ident["content_key"]:
            print("already in library as {} ({})".format(c["id"], c["library_path"]))
            return 0
    fields = {"category": args.category, "description": args.description, "tags": args.tags or [],
              "people": args.people, "status": args.status, "used_in": args.used_in or [],
              "model": args.model, "prompt": args.prompt, "notes": args.notes or ""}
    rec = new_record(cat, src, args.kind, args.slug, fields)
    place(src, rec, move=False)
    make_thumbs(rec)
    cat["clips"].append(rec)
    save(cat)
    print("added {} -> {}".format(rec["id"], LIB / rec["library_path"]))
    if not args.no_drive:
        drive_sync(cat, only=[rec["id"]])
        save(cat)
        print("drive:", drive_link(rec))
    print("Run `sheet` to refresh the Google Sheet.")
    return 0


def tokens(text):
    return re.findall(r"[a-z0-9]+", str(text).lower())


def score(clip, q):
    fields = {"slug": 3, "description": 2, "tags": 3, "category": 2, "prompt": 1, "notes": 1, "legacy_name": 1}
    total = 0
    for word in q:
        best = 0
        stem = word[:-1] if word.endswith("s") and len(word) > 3 else word
        for f, weight in fields.items():
            val = clip.get(f)
            val = " ".join(val) if isinstance(val, list) else (val or "")
            toks = tokens(val)
            if any(t == word or t.startswith(stem) for t in toks):
                best = max(best, weight)
        if best == 0:
            return 0
        total += best
    return total


def cmd_find(args):
    cat = load()
    q = tokens(args.words)
    hits = []
    for c in cat["clips"]:
        if args.kind and c["kind"] != args.kind:
            continue
        if args.aspect and c["aspect"] != args.aspect:
            continue
        if args.people and c["people"] != args.people:
            continue
        if args.category and c["category"] != args.category:
            continue
        if args.min_seconds and c["duration"] < args.min_seconds:
            continue
        s = score(c, q) if q else 1
        if s:
            hits.append((s, c))
    hits.sort(key=lambda x: (-x[0], x[1]["id"]))
    for s, c in hits[: args.limit]:
        mounted = (LIB / c["library_path"]).exists()
        print("{} | {} | {} | {:.1f}s | {} | {} | used: {}".format(
            c["id"], c["kind"], c["aspect"], c["duration"], c["people"], c["status"], ", ".join(c["used_in"]) or "-"))
        print("    {}".format(c["description"]))
        print("    {}  {}".format(LIB / c["library_path"] if mounted else "(Extreme unplugged) " + c["library_path"], drive_link(c) or ""))
        print("    preview: {}".format(CONTACT / (c["id"] + ".jpg")))
    if not hits:
        print("No library matches.")
    if args.rolls:
        print("\n--- raw shoot rolls (roll_sidecar find) ---", flush=True)
        subprocess.run([sys.executable, str(HERE.parent / "rolls" / "roll_sidecar.py"), "find", args.words])
    return 0 if hits else 1


def cmd_show(args):
    cat = load()
    for c in cat["clips"]:
        if c["id"].lower() == args.id.lower():
            print(json.dumps(c, indent=1, ensure_ascii=False))
            return 0
    print("not found")
    return 1


def cmd_thumbs(args):
    cat = load()
    for c in cat["clips"]:
        if args.force or not (CONTACT / (c["id"] + ".jpg")).exists():
            make_thumbs(c)
    print("ok")


DESCRIBE_PROMPT = (
    "Six frames sampled evenly from one short video clip that will be filed in a stock-footage library "
    "for video editors. Return JSON only: {\"description\": one concrete sentence (who, doing what, where, "
    "camera move), \"tags\": 5-10 lowercase search words, \"people\": one of dan-like|other-man|woman|mixed|none, "
    "\"defects\": list of visible AI-generation defects (melting or extra fingers, warped face or body, garbled "
    "text, objects morphing, smoke or fog that should not be there) or [] if none, \"burned_text\": any on-screen "
    "text or watermark you can read, else \"\"}. The editor's draft description is: ")


def cmd_describe(args):
    cat = load()
    key = rs.secret_value("GEMINI_API_KEY")
    model = args.model
    spent = 0.0
    todo = [c for c in cat["clips"] if args.force or "vision" not in c]
    if args.ids:
        todo = [c for c in todo if c["id"] in args.ids]
    print("describing {} clips, est ${:.2f}".format(len(todo), len(todo) * 0.0015), flush=True)
    for c in todo:
        img = CONTACT / (c["id"] + ".jpg")
        if not img.exists():
            continue
        payload = {"contents": [{"parts": [{"text": DESCRIBE_PROMPT + json.dumps(c.get("description", ""))},
                                           {"inline_data": {"mime_type": "image/jpeg", "data": base64.b64encode(img.read_bytes()).decode()}}]}],
                   "generationConfig": {"temperature": 0.1, "maxOutputTokens": 800, "responseMimeType": "application/json"}}
        url = "https://generativelanguage.googleapis.com/v1beta/models/{}:generateContent?key={}".format(model, key)
        req = urllib.request.Request(url, data=json.dumps(payload).encode(), headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                body = json.loads(r.read())
            usage = rs.description_token_usage(body)
            spent += usage["estimated_actual_usd"]
            text = "".join(p.get("text", "") for cand in body.get("candidates", []) for p in cand.get("content", {}).get("parts", []))
            parsed = rs.parse_json_text(text)
            c["vision"] = parsed[0] if isinstance(parsed, list) and parsed else parsed
            c["vision"]["model"] = model
        except Exception as exc:  # keep going; one bad clip must not stop the batch
            c["vision"] = {"error": str(exc)[:200]}
        for t in c.get("vision", {}).get("tags", []) or []:
            if t not in c["tags"]:
                c["tags"].append(t)
    save(cat)
    print("spent about ${:.3f}".format(spent))


# ---------------------------------------------------------------- Drive + Sheet

def drive_token():
    subprocess.run([RCLONE, "about", "gdrive:"], capture_output=True)  # refreshes the token if stale
    dump = json.loads(subprocess.run([RCLONE, "config", "dump"], capture_output=True, text=True).stdout)
    return json.loads(dump["gdrive"]["token"])["access_token"]


def api(method, url, token, body=None, ctype="application/json"):
    data = json.dumps(body).encode() if isinstance(body, (dict, list)) else body
    req = urllib.request.Request(url, data=data, method=method,
                                 headers={"Authorization": "Bearer " + token, "Content-Type": ctype})
    for attempt in range(6):
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                raw = r.read()
                return json.loads(raw) if raw else {}
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode(errors="replace")
            if exc.code in (403, 429, 500, 503) and ("Quota" in detail or "rate" in detail.lower() or exc.code >= 429):
                time.sleep(6 * (attempt + 1))
                continue
            raise RuntimeError("{} {} -> {} {}".format(method, url, exc.code, detail[:500]))
    raise RuntimeError("gave up after retries: " + url)


def drive_link(c):
    return "https://drive.google.com/file/d/{}/view".format(c["drive_id"]) if c.get("drive_id") else ""


def drive_folder_ids(token, rel_folder):
    """Resolve (creating as needed) the Drive folder for a library-relative folder path."""
    parent = DRIVE_ROOT
    for part in Path(rel_folder).parts:
        q = "'{}' in parents and name = '{}' and mimeType = 'application/vnd.google-apps.folder' and trashed = false".format(
            parent, part.replace("'", "\\'"))
        res = api("GET", "https://www.googleapis.com/drive/v3/files?fields=files(id)&q=" + urllib.parse.quote(q), token)
        if res.get("files"):
            parent = res["files"][0]["id"]
        else:
            parent = api("POST", "https://www.googleapis.com/drive/v3/files?fields=id", token,
                         {"name": part, "mimeType": "application/vnd.google-apps.folder", "parents": [parent]})["id"]
    return parent


RCLONE_EXCLUDES = ["--exclude", "._*", "--exclude", ".DS_Store", "--exclude", "08 SixPackAbs Archive*/**"]


def drive_sync(cat, only=None):
    """Mirror the library to Drive (rclone creates folders), then record each clip's Drive file id.

    08 SixPackAbs Archive stays off Drive: it is quarantined and the Drive folder is anyone-with-link."""
    if only:
        for c in cat["clips"]:
            if c["id"] in only and (LIB / c["library_path"]).exists():
                subprocess.run([RCLONE, "copyto", str(LIB / c["library_path"]), "gdrive:" + c["library_path"],
                                "--drive-root-folder-id=" + DRIVE_ROOT, "-q"], check=True)
    else:
        subprocess.run([RCLONE, "copy", str(LIB), "gdrive:", "--drive-root-folder-id=" + DRIVE_ROOT,
                        "--transfers", "4", "--drive-chunk-size", "64M", "--stats", "60s", "--stats-one-line",
                        "-v"] + RCLONE_EXCLUDES, check=True)
    for top in (AI_DIR, BROLL_DIR):
        res = subprocess.run([RCLONE, "lsjson", "-R", "--files-only", "gdrive:" + top, "--drive-root-folder-id=" + DRIVE_ROOT],
                             capture_output=True, text=True, check=True)
        ids = {top + "/" + r["Path"]: r["ID"] for r in json.loads(res.stdout)}
        for c in cat["clips"]:
            if c["library_path"] in ids:
                c["drive_id"] = ids[c["library_path"]]
    return cat


def cmd_drive_sync(args):
    cat = load()
    try:
        drive_sync(cat)
    finally:
        save(cat)
    print("ok")


SHEET_COLS = ["ID", "Preview", "Description", "Kind", "Category", "People", "Shape", "Seconds", "Sound",
              "Status", "Used in", "Tags", "AI label needed", "Notes", "Drive link", "File name"]


def sheet_rows(cat):
    rows = [SHEET_COLS]
    for c in cat["clips"]:
        prev = '=IMAGE("https://drive.google.com/thumbnail?id={}&sz=w320")'.format(c["drive_id"]) if c.get("drive_id") else ""
        rows.append([c["id"], prev, c["description"], c["kind"], c["category"], c["people"], c["aspect"],
                     round(c["duration"], 1), "yes" if c["has_audio"] else "no",
                     "used in a finished video" if c["status"] == "used-final" else "not used yet",
                     ", ".join(c["used_in"]), ", ".join(c["tags"][:12]), "YES" if c.get("needs_ai_label") else "",
                     c.get("notes", ""), drive_link(c), c["file_name"]])
    return rows


def build_xlsx(cat, path):
    """Formatted workbook; Drive converts it to a Google Sheet (Sheets API is not enabled on rclone's project)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Clips"
    rows = sheet_rows(cat)
    for r in rows:
        ws.append(r)
    widths = [8, 24, 60, 9, 18, 11, 9, 9, 8, 20, 26, 36, 11, 36, 36, 44]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F5")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i in range(2, len(rows) + 1):
        ws.row_dimensions[i].height = 84
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:{}{}".format(chr(64 + len(SHEET_COLS)), len(rows))
    howto = wb.create_sheet("How to use")
    for line in HOWTO:
        howto.append([line])
    howto.column_dimensions["A"].width = 120
    wb.save(path)
    return len(rows) - 1


HOWTO = [
    "CLIP LIBRARY: every reusable AI clip and B-roll clip we own. Check here before making or buying a new clip.",
    "",
    "Find a clip: use the filter arrows on the Clips tab (Kind, Category, People, Shape), or Ctrl/Cmd+F for a word like salad, sleep, gym, phone.",
    "Open it: click the Drive link. Every file is view-for-anyone-with-the-link and downloadable.",
    "The ID (A0042 = AI clip, B0031 = B-roll) is the clip's permanent name. Quote the ID in revision notes.",
    "Status 'used in a finished video' means it is already in the videos listed under Used in: avoid repeating it in the same audience's feed.",
    "AI label needed = YES: burn the on-screen 'AI-GENERATED' tag while the clip is on screen.",
    "Before/after rule: a before and its after must be the same person. Never pair two different people.",
    "Folders: 04 AI-Generated Clips/<category>, 03 B-Roll - Real Footage/<dan-filmed | screen-recordings | stock>.",
    "This sheet is rebuilt from the catalog; edits typed here are overwritten. Send corrections to Dan or Claude.",
]


def cmd_sheet(args):
    cat = load()
    token = drive_token()
    xlsx = CAT_DIR / "clip-library.xlsx"
    n = build_xlsx(cat, xlsx)
    data = xlsx.read_bytes()
    xmime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    sid = cat.get("sheet_id")
    if sid:
        api("PATCH", "https://www.googleapis.com/upload/drive/v3/files/{}?uploadType=media".format(sid), token, data, ctype=xmime)
    else:
        boundary = "cliplibboundary"
        meta = json.dumps({"name": "Clip Library - AI clips and B-roll (catalog)",
                           "mimeType": "application/vnd.google-apps.spreadsheet", "parents": [DRIVE_ROOT]})
        body = ("--{b}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{m}\r\n--{b}\r\nContent-Type: {x}\r\n\r\n"
                .format(b=boundary, m=meta, x=xmime)).encode() + data + "\r\n--{}--\r\n".format(boundary).encode()
        res = api("POST", "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&fields=id", token, body,
                  ctype="multipart/related; boundary=" + boundary)
        sid = cat["sheet_id"] = res["id"]
        save(cat)
    print("https://docs.google.com/spreadsheets/d/{}/edit  ({} clips)".format(sid, n))


def cmd_verify(args):
    cat = load()
    bad = 0
    ids = set()
    for c in cat["clips"]:
        if c["id"] in ids:
            print("DUPLICATE ID", c["id"]); bad += 1
        ids.add(c["id"])
        if LIB.exists() and not (LIB / c["library_path"]).exists():
            print("MISSING FILE", c["id"], c["library_path"]); bad += 1
        if not (CONTACT / (c["id"] + ".jpg")).exists():
            print("NO PREVIEW", c["id"]); bad += 1
        if not c.get("description"):
            print("NO DESCRIPTION", c["id"]); bad += 1
        if not c.get("drive_id"):
            print("NOT ON DRIVE", c["id"])
    print("{} clips, {} problems".format(len(cat["clips"]), bad))
    return 1 if bad else 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    f = sub.add_parser("find", help="search the catalog (works with the Extreme drive unplugged)")
    f.add_argument("words")
    f.add_argument("--kind", choices=KINDS)
    f.add_argument("--aspect", help="9x16, 16x9, 1x1 ...")
    f.add_argument("--people", choices=PEOPLE)
    f.add_argument("--category")
    f.add_argument("--min-seconds", type=float)
    f.add_argument("--limit", type=int, default=15)
    f.add_argument("--rolls", action="store_true", help="also search raw shoot rolls via roll_sidecar")
    s = sub.add_parser("show"); s.add_argument("id")
    a = sub.add_parser("add", help="register a newly approved clip (copies it in, previews, Drive)")
    a.add_argument("file")
    a.add_argument("--kind", choices=KINDS, required=True)
    a.add_argument("--slug", required=True)
    a.add_argument("--description", required=True)
    a.add_argument("--category", default="")
    a.add_argument("--people", choices=PEOPLE, default="none")
    a.add_argument("--status", choices=STATUSES, default="usable-unused")
    a.add_argument("--used-in", nargs="*")
    a.add_argument("--tags", nargs="*")
    a.add_argument("--model"); a.add_argument("--prompt"); a.add_argument("--notes")
    a.add_argument("--no-drive", action="store_true")
    i = sub.add_parser("ingest"); i.add_argument("files", nargs="+"); i.add_argument("--dry-run", action="store_true")
    t = sub.add_parser("thumbs"); t.add_argument("--force", action="store_true")
    d = sub.add_parser("describe"); d.add_argument("--force", action="store_true"); d.add_argument("--ids", nargs="*")
    d.add_argument("--model", default=rs.DEFAULT_MODEL)
    sub.add_parser("drive-sync")
    sub.add_parser("sheet")
    sub.add_parser("verify")
    args = ap.parse_args()
    fn = {"find": cmd_find, "show": cmd_show, "add": cmd_add, "ingest": cmd_ingest, "thumbs": cmd_thumbs,
          "describe": cmd_describe, "drive-sync": cmd_drive_sync, "sheet": cmd_sheet, "verify": cmd_verify}[args.cmd]
    sys.exit(fn(args) or 0)


if __name__ == "__main__":
    main()
